from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Annotated
import secrets
import json
import aiohttp
from fastapi import FastAPI, Depends, HTTPException, Header, UploadFile, File, Form
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession
from openpyxl import load_workbook, Workbook
from .db import Session, init_db, settings, Organization, Employee, Survey, Candidate, Channel, Vote

app = FastAPI(title='Survey Bot Admin API', docs_url=None, redoc_url=None, openapi_url=None)

@app.on_event('startup')
async def startup():
    await init_db()

async def auth(x_admin_key: Annotated[str | None, Header()] = None):
    if not x_admin_key or not secrets.compare_digest(x_admin_key, settings.admin_api_key):
        raise HTTPException(401, 'Admin kaliti noto‘g‘ri')

async def session():
    async with Session() as db:
        yield db

class OrgIn(BaseModel):
    name: str = Field(min_length=2, max_length=200)

class SurveyIn(BaseModel):
    organization_id: int
    title: str = Field(min_length=2, max_length=250)
    description: str = ''
    ends_at: datetime

class CandidateIn(BaseModel):
    employee_id: int

class ChannelIn(BaseModel):
    chat_id: int
    invite_url: str

@app.get('/')
async def home():
    return FileResponse(Path(__file__).parent / 'static' / 'index.html')

@app.post('/api/organizations', dependencies=[Depends(auth)])
async def add_org(data: OrgIn, db: AsyncSession = Depends(session)):
    org = Organization(name=data.name)
    db.add(org)
    await db.commit()
    await db.refresh(org)
    return {'id':org.id,'name':org.name}

@app.get('/api/organizations', dependencies=[Depends(auth)])
async def list_orgs(db: AsyncSession = Depends(session)):
    rows = (await db.scalars(select(Organization).order_by(Organization.id))).all()
    return [{'id':o.id,'name':o.name} for o in rows]

@app.post('/api/surveys', dependencies=[Depends(auth)])
async def add_survey(data: SurveyIn, db: AsyncSession = Depends(session)):
    if not await db.get(Organization, data.organization_id):
        raise HTTPException(404,'Tashkilot topilmadi')
    if data.ends_at.tzinfo is None or data.ends_at <= datetime.now(timezone.utc):
        raise HTTPException(400,'Kelajakdagi vaqt va vaqt zonasi kerak')
    survey = Survey(**data.model_dump())
    db.add(survey)
    await db.commit()
    await db.refresh(survey)
    return {'id':survey.id,'title':survey.title,'bot_link':f'https://t.me/{settings.public_bot_username}?start=s{survey.id}'}

@app.get('/api/surveys', dependencies=[Depends(auth)])
async def list_surveys(db: AsyncSession = Depends(session)):
    rows = (await db.scalars(select(Survey).order_by(Survey.id.desc()))).all()
    return [{'id':s.id,'organization_id':s.organization_id,'title':s.title,'ends_at':s.ends_at.isoformat(),'active':s.active,'bot_link':f'https://t.me/{settings.public_bot_username}?start=s{s.id}'} for s in rows]

@app.post('/api/surveys/{survey_id}/candidates', dependencies=[Depends(auth)])
async def add_candidate(survey_id:int, data:CandidateIn, db:AsyncSession=Depends(session)):
    survey, employee = await db.get(Survey,survey_id), await db.get(Employee,data.employee_id)
    if not survey or not employee or survey.organization_id != employee.organization_id:
        raise HTTPException(400,'Nomzod va so‘rovnoma tashkiloti mos emas')
    row = Candidate(survey_id=survey_id,employee_id=data.employee_id)
    db.add(row)
    await db.commit()
    return {'id':row.id}

@app.post('/api/surveys/{survey_id}/channels', dependencies=[Depends(auth)])
async def add_channel(survey_id:int, data:ChannelIn, db:AsyncSession=Depends(session)):
    if not await db.get(Survey,survey_id):
        raise HTTPException(404,'So‘rovnoma topilmadi')
    if not data.invite_url.startswith('https://t.me/'):
        raise HTTPException(400,'Telegram kanal havolasi kerak')
    row = Channel(survey_id=survey_id,**data.model_dump())
    db.add(row)
    await db.commit()
    return {'id':row.id}

@app.post('/api/organizations/{org_id}/employees/import', dependencies=[Depends(auth)])
async def import_employees(org_id:int, file:UploadFile=File(...), db:AsyncSession=Depends(session)):
    if not await db.get(Organization,org_id):
        raise HTTPException(404,'Tashkilot topilmadi')
    raw = await file.read(5_000_001)
    if len(raw)>5_000_000 or not (file.filename or '').lower().endswith('.xlsx'):
        raise HTTPException(400,'Faqat 5 MB gacha .xlsx')
    try:
        wb = load_workbook(BytesIO(raw),read_only=True,data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = [str(x or '').strip().lower() for x in next(rows)]
        required = ['id','name','region','position']
        if not all(c in header for c in required):
            raise ValueError('Ustunlar: id, name, region, position')
        idx = [header.index(c) for c in required]
        parsed = []
        for row in rows:
            if not any(v is not None for v in row): continue
            vals = [str(row[i] if i<len(row) and row[i] is not None else '').strip() for i in idx]
            if not vals[0] or not vals[1]: raise ValueError('ID va name bo‘sh bo‘lmasin')
            parsed.append(vals)
            if len(parsed)>10000: raise ValueError('Bir yuklashda ko‘pi bilan 10000 xodim')
        if len({r[0] for r in parsed})!=len(parsed): raise ValueError('Takroriy ID')
    except Exception as exc:
        raise HTTPException(400,f'Excel xatosi: {exc}')
    existing = {e.external_id:e for e in (await db.scalars(select(Employee).where(Employee.organization_id==org_id))).all()}
    created=updated=0
    for ext,name,region,position in parsed:
        if ext in existing:
            e=existing[ext]
            e.name,e.region,e.position,e.active=name,region,position,True
            updated+=1
        else:
            db.add(Employee(organization_id=org_id,external_id=ext,name=name,region=region,position=position))
            created+=1
    await db.commit()
    return {'created':created,'updated':updated,'unchanged_other_employees':True}

@app.get('/api/organizations/{org_id}/employees', dependencies=[Depends(auth)])
async def employees(org_id:int, db:AsyncSession=Depends(session)):
    rows=(await db.scalars(select(Employee).where(Employee.organization_id==org_id).order_by(Employee.id).limit(10000))).all()
    return [{'id':e.id,'external_id':e.external_id,'name':e.name,'region':e.region,'position':e.position} for e in rows]

@app.get('/api/surveys/{survey_id}/results', dependencies=[Depends(auth)])
async def results(survey_id:int, db:AsyncSession=Depends(session)):
    q=(select(Candidate.id,Employee.name,Employee.region,func.count(Vote.id).label('votes'))
       .join(Employee,Employee.id==Candidate.employee_id)
       .outerjoin(Vote,Vote.candidate_id==Candidate.id)
       .where(Candidate.survey_id==survey_id)
       .group_by(Candidate.id,Employee.name,Employee.region)
       .order_by(func.count(Vote.id).desc()))
    rows=(await db.execute(q)).all()
    return [{'candidate_id':r.id,'name':r.name,'region':r.region,'votes':r.votes} for r in rows]

@app.get('/api/surveys/{survey_id}/results.xlsx', dependencies=[Depends(auth)])
async def export_results(survey_id:int, db:AsyncSession=Depends(session)):
    data=await results(survey_id,db)
    wb=Workbook();ws=wb.active;ws.title='Natijalar'
    ws.append(['Nomzod ID','F.I.Sh.','Hudud','Ovozlar'])
    for r in data: ws.append([r['candidate_id'],r['name'],r['region'],r['votes']])
    output=BytesIO();wb.save(output);output.seek(0)
    return StreamingResponse(output,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="survey_{survey_id}.xlsx"'})


@app.post('/api/surveys/{survey_id}/publish-photo', dependencies=[Depends(auth)])
async def publish_survey_photo(
    survey_id: int,
    chat_id: int = Form(...),
    caption: str = Form(...),
    photo: UploadFile = File(...),
    db: AsyncSession = Depends(session),
):
    """Publish one survey announcement to a channel registered for that survey."""
    survey = await db.get(Survey, survey_id)
    if survey is None:
        raise HTTPException(404, 'So‘rovnoma topilmadi')
    allowed_channel = await db.scalar(
        select(Channel).where(Channel.survey_id == survey_id, Channel.chat_id == chat_id)
    )
    if allowed_channel is None:
        raise HTTPException(400, 'Kanal bu so‘rovnomaga biriktirilmagan')
    if not survey.active:
        raise HTTPException(400, 'So‘rovnoma faol emas')
    if not caption.strip() or len(caption) > 1024:
        raise HTTPException(400, 'Post matni 1–1024 belgi bo‘lishi kerak')

    content = await photo.read(8_000_001)
    if len(content) > 8_000_000:
        raise HTTPException(400, 'Rasm hajmi 8 MB dan oshmasin')
    is_jpeg = content.startswith(b'\xff\xd8\xff')
    is_png = content.startswith(b'\x89PNG\r\n\x1a\n')
    if not (is_jpeg or is_png):
        raise HTTPException(400, 'Faqat haqiqiy JPG yoki PNG rasm qabul qilinadi')

    username = settings.public_bot_username.lstrip('@')
    vote_url = f'https://t.me/{username}?start=s{survey_id}'
    keyboard = {'inline_keyboard': [[{'text': '🗳 OVOZ BERISH', 'url': vote_url}]]}
    form = aiohttp.FormData()
    form.add_field('chat_id', str(chat_id))
    form.add_field('caption', caption.strip())
    form.add_field('reply_markup', json.dumps(keyboard, ensure_ascii=False))
    form.add_field('photo', content, filename='post.jpg' if is_jpeg else 'post.png',
                   content_type='image/jpeg' if is_jpeg else 'image/png')
    url = f'https://api.telegram.org/bot{settings.bot_token}/sendPhoto'
    try:
        timeout = aiohttp.ClientTimeout(total=35)
        async with aiohttp.ClientSession(timeout=timeout) as client:
            async with client.post(url, data=form) as response:
                payload = await response.json(content_type=None)
    except (aiohttp.ClientError, TimeoutError, ValueError):
        raise HTTPException(502, 'Telegram bilan bog‘lanishda xatolik')
    if not payload.get('ok'):
        # Telegram error text contains no token, but do not expose raw response.
        raise HTTPException(502, 'Post yuborilmadi. Botning kanal administratorligi va post joylash huquqini tekshiring')
    return {
        'ok': True,
        'survey_id': survey_id,
        'chat_id': chat_id,
        'message_id': payload['result']['message_id'],
        'bot_link': vote_url,
    }

# So'rovnomani faolsizlantirish
@app.post(
    "/api/surveys/{survey_id}/deactivate",
    dependencies=[Depends(auth)]
)
async def deactivate_survey(
    survey_id: int,
    db: AsyncSession = Depends(session)
):
    survey = await db.get(Survey, survey_id)
    if not survey:
        raise HTTPException(404, "So'rovnoma topilmadi")

    survey.active = False
    await db.commit()

    return {
        "survey_id": survey_id,
        "active": False,
        "message": "So'rovnoma faolsizlantirildi"
    }


# Majburiy kanallar ro'yxati
@app.get(
    "/api/surveys/{survey_id}/channels",
    dependencies=[Depends(auth)]
)
async def list_channels(
    survey_id: int,
    db: AsyncSession = Depends(session)
):
    if not await db.get(Survey, survey_id):
        raise HTTPException(404, "So'rovnoma topilmadi")

    rows = (
        await db.scalars(
            select(Channel)
            .where(Channel.survey_id == survey_id)
            .order_by(Channel.id)
        )
    ).all()

    return [
        {
            "id": ch.id,
            "chat_id": ch.chat_id,
            "invite_url": ch.invite_url
        }
        for ch in rows
    ]


# Kanalni majburiy obunadan olib tashlash
@app.delete(
    "/api/surveys/{survey_id}/channels/{channel_id}",
    dependencies=[Depends(auth)]
)
async def remove_channel(
    survey_id: int,
    channel_id: int,
    db: AsyncSession = Depends(session)
):
    channel = await db.get(Channel, channel_id)

    if not channel or channel.survey_id != survey_id:
        raise HTTPException(404, "Kanal topilmadi")

    await db.delete(channel)
    await db.commit()

    return {
        "ok": True,
        "message": "Kanal majburiy obunadan olib tashlandi"
    }
